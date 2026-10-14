import asyncio
import json
from io import BytesIO
from uuid import UUID, uuid4

import httpx
import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import func, select
from sqlalchemy.dialects import postgresql
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker

from app.core.config import Settings
from app.core.errors import AppError
from app.db.base import Base
from app.db.session import make_engine
from app.models import domain as m
from app.services import factories
from app.services.inference.mock import MockModelProvider
from app.services.knowledge.chunking import chunk_text
from app.services.knowledge.embeddings import (
    MockEmbeddingProvider,
    OpenAIEmbeddingProvider,
    normalize,
)
from app.services.knowledge.extraction import extract
from app.services.knowledge.mock import MockKnowledgeService
from app.services.knowledge.progress import observer
from app.services.knowledge.questoriaa import QuestoriaaKnowledgeService
from app.services.knowledge.schemas import CourseDraft
from app.services.knowledge.store import DatabaseChunkStore, vector_search_statement
from app.services.storage.local import LocalStorageService
from app.workers.processing import execute_job, run_once


def run(awaitable):
    return asyncio.run(awaitable)


def pdf_bytes(texts=("Plants use sunlight to produce sugar.", "Roots absorb water from soil.")):
    with pymupdf.open() as document:
        for text in texts:
            document.new_page().insert_text((72, 72), text)
        return document.tobytes()


def pptx_bytes():
    presentation = Presentation()
    for index in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(7), Inches(1)
        ).text = f"Slide topic {index + 1}"
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(7), Inches(1)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "Quantity", "10 units"
        slide.notes_slide.notes_text_frame.text = f"Explain the quantity on slide {index + 1}."
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


class OutlineModel(MockModelProvider):
    async def generate_structured(self, messages, schema):
        if schema is not CourseDraft:
            return await super().generate_structured(messages, schema)
        payload = json.loads(messages[-1]["content"])
        assert len(payload["sources"]) == 2
        assert "page_number" not in payload["sources"][0]
        return CourseDraft.model_validate(
            {
                "chapters": [
                    {
                        "title": "Botany",
                        "topics": [
                            {
                                "title": "Plant systems",
                                "concepts": [
                                    {
                                        "name": "Photosynthesis",
                                        "description": "Plants produce sugar.",
                                        "source_numbers": [1],
                                    },
                                    {
                                        "name": "Root uptake",
                                        "description": "Roots absorb water.",
                                        "difficulty": 2,
                                        "source_numbers": [2],
                                        "prerequisite_names": ["Photosynthesis"],
                                    },
                                    {
                                        "name": "photo synthesis",
                                        "description": "A duplicate.",
                                        "source_numbers": [1],
                                    },
                                ],
                            }
                        ],
                    }
                ]
            }
        )


@pytest.fixture
def uploaded(course, factory, tmp_path):
    storage = LocalStorageService(tmp_path / "objects", "/api/v1")
    data = pdf_bytes()
    key = f"{course['id']}/notes.pdf"
    storage.path(key).parent.mkdir(parents=True)
    storage.path(key).write_bytes(data)
    with factory() as db:
        file = m.CourseFile(
            course_id=UUID(course["id"]),
            filename="notes.pdf",
            source_type="pdf",
            object_key=key,
            content_type="application/pdf",
            size_bytes=len(data),
            upload_completed=True,
        )
        db.add(file)
        db.commit()
        return UUID(course["id"]), file.id, storage


def service(db, storage, model=None, embeddings=None):
    return QuestoriaaKnowledgeService(
        db, storage, model or OutlineModel(), embeddings or MockEmbeddingProvider()
    )


def test_pdf_locations_and_cleaning():
    chunks = extract(pdf_bytes(), "pdf")
    assert [c.page_number for c in chunks] == [1, 2]
    assert "sunlight" in chunks[0].content and "Roots" in chunks[1].content
    assert all(c.slide_number is None for c in chunks)


def test_pptx_locations_tables_and_notes():
    chunks = extract(pptx_bytes(), "pptx")
    assert [c.slide_number for c in chunks] == [1, 2]
    assert all("Quantity | 10 units" in c.content and "Speaker notes:" in c.content for c in chunks)
    assert "slide 2" not in chunks[0].content


@pytest.mark.parametrize("source_type", ["image", "video", "text"])
def test_unsupported_source_types(source_type):
    with pytest.raises(AppError) as error:
        extract(b"content", source_type)
    assert error.value.code == "UNSUPPORTED_INGESTION"


@pytest.mark.parametrize("source_type", ["pdf", "pptx"])
def test_malformed_document_is_controlled(source_type):
    with pytest.raises(AppError) as error:
        extract(b"not a document", source_type)
    assert error.value.code == "INVALID_DOCUMENT"


def test_blank_pdf_page_is_skipped_without_renumbering():
    chunks = extract(pdf_bytes(["Readable page one", "", "Readable page three"]), "pdf")
    assert [chunk.page_number for chunk in chunks] == [1, 3]


def test_empty_pptx_slide_is_skipped_without_renumbering():
    presentation = Presentation()
    for text in ["First slide", "", "Third slide"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if text:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = text
    output = BytesIO()
    presentation.save(output)
    assert [chunk.slide_number for chunk in extract(output.getvalue(), "pptx")] == [1, 3]


@pytest.mark.parametrize("source_type", ["pdf", "pptx"])
def test_entirely_textless_document_fails(source_type):
    if source_type == "pdf":
        data = pdf_bytes(["", ""])
    else:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        # An empty table must not manufacture extractable separator characters.
        slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(7), Inches(1))
        presentation.slides.add_slide(presentation.slide_layouts[6])
        output = BytesIO()
        presentation.save(output)
        data = output.getvalue()
    with pytest.raises(AppError) as error:
        extract(data, source_type)
    assert error.value.code == "TEXTLESS_DOCUMENT"


def test_chunking_bounds_and_code():
    chunks = chunk_text("Introduction:\nint x = 2;\nreturn x;\n\n" + "longword" * 500)
    assert all(0 < len(c) <= 1400 for c in chunks)
    assert "int x = 2;\nreturn x;" in chunks[0]


def test_embeddings_are_deterministic_and_normalized():
    texts = ["plants sunlight", "water roots"]
    first = run(MockEmbeddingProvider().embed_texts(texts))
    assert first == run(MockEmbeddingProvider().embed_texts(texts))
    assert first[0] == run(MockEmbeddingProvider().embed_query(texts[0]))
    assert first[0] != first[1]
    assert all(len(v) == 1024 and sum(x * x for x in v) == pytest.approx(1) for v in first)


@pytest.mark.parametrize(
    "vector", [[1.0], [0.0] * 1024, [float("nan")] * 1024, [float("inf")] * 1024]
)
def test_invalid_embeddings_rejected(vector):
    with pytest.raises(ValueError):
        normalize(vector)


def test_worker_persists_graph_and_retrieval_across_sessions(
    client, uploaded, factory, monkeypatch
):
    course_id, file_id, storage = uploaded
    settings = Settings(_env_file=None, knowledge_provider="questoriaa", embedding_provider="mock")
    monkeypatch.setattr(factories, "get_settings", lambda: settings)
    monkeypatch.setattr(factories, "get_storage", lambda: storage)
    monkeypatch.setattr(factories, "create_model", OutlineModel)
    job_id = client.post(f"/api/v1/courses/{course_id}/process").json()["id"]
    assert run(run_once(factory))
    assert client.get(f"/api/v1/jobs/{job_id}").json()["status"] == "completed"
    # API creates a new service and session using the same central factory.
    mapping = client.get(f"/api/v1/courses/{course_id}/map").json()
    concepts = mapping["chapters"][0]["topics"][0]["concepts"]
    assert len(concepts) == 2
    assert len(mapping["prerequisites"]) == 1
    ids = {c["name"]: c["id"] for c in concepts}
    link = mapping["prerequisites"][0]
    assert link["concept_id"] == ids["Root uptake"]
    assert link["prerequisite_concept_id"] == ids["Photosynthesis"]
    with factory() as db:
        fresh = factories.create_knowledge(db)
        result = run(fresh.retrieve_context(course_id, "Roots water soil"))
        assert result[0].source.page_number == 2 and result[0].source.file_id == file_id
        assert result[0].score > result[1].score
        filtered = run(fresh.retrieve_context(course_id, "water", UUID(ids["Photosynthesis"])))
        assert len(filtered) == 1 and filtered[0].source.page_number == 1
        assert run(fresh.retrieve_context(uuid4(), "water")) == []
        assert run(fresh.retrieve_context(course_id, "water", uuid4())) == []
        assert run(fresh.retrieve_context(course_id, "water", top_k=0)) == []
        assert all(len(row.embedding) == 1024 for row in db.scalars(select(m.SourceChunk)))
    # Existing tutor consumes retrieved metadata without a parallel tutor stack.
    response = client.post(
        "/api/v1/tutor/chat", json={"course_id": str(course_id), "message": "water"}
    )
    assert response.status_code == 200, response.text
    assert {source["file_id"] for source in response.json()["sources"]} == {str(file_id)}
    assert {source["page_number"] for source in response.json()["sources"]} == {1, 2}
    assert "Roots absorb water" in response.json()["answer"]


@pytest.mark.parametrize("mock", [False, True])
def test_exact_ids_order_scope_and_authoritative_metadata(uploaded, factory, mock):
    course_id, _, storage = uploaded
    with factory() as db:
        run(service(db, storage).process_course(course_id))
        db.commit()
    with factory() as db:
        knowledge = MockKnowledgeService(db) if mock else service(db, storage)
        rows = db.scalars(select(m.SourceChunk).order_by(m.SourceChunk.page_number)).all()
        requested = [rows[1].id, uuid4(), rows[0].id, rows[1].id]
        found = run(knowledge.get_context_by_ids(course_id, requested))
        assert [c.chunk_id for c in found] == [rows[1].id, rows[0].id]
        assert [c.source.page_number for c in found] == [2, 1]
        assert run(knowledge.get_context_by_ids(uuid4(), requested)) == []
        scoped = run(knowledge.get_context_by_ids(course_id, requested, rows[0].concept_id))
        assert [c.chunk_id for c in scoped] == [rows[0].id]
        assert run(knowledge.get_context_by_ids(course_id, [])) == []


def test_reprocessing_rejected_without_changes(uploaded, factory):
    course_id, _, storage = uploaded
    with factory() as db:
        knowledge = service(db, storage)
        run(knowledge.process_course(course_id))
        db.commit()
        original_ids = list(db.scalars(select(m.SourceChunk.id)))
        with pytest.raises(AppError) as error:
            run(knowledge.process_course(course_id))
        assert error.value.code == "ALREADY_PROCESSED" and error.value.status == 409
        db.commit()
        assert list(db.scalars(select(m.SourceChunk.id))) == original_ids
        assert db.scalar(select(func.count()).select_from(m.Concept)) == 2


def test_failure_rolls_back_graph_and_fails_job(uploaded, factory):
    class FailedEmbedding(MockEmbeddingProvider):
        async def embed_texts(self, texts):
            raise AppError(503, "EMBEDDING_UNAVAILABLE", "Embedding endpoint unavailable.")

    course_id, file_id, storage = uploaded
    with factory() as db:
        job = m.ProcessingJob(course_id=course_id, file_id=file_id)
        db.add(job)
        db.commit()
        run(execute_job(db, job.id, service(db, storage, embeddings=FailedEmbedding())))
        assert db.get(m.ProcessingJob, job.id).status == "failed"
        assert db.get(m.ProcessingJob, job.id).error.startswith("EMBEDDING_UNAVAILABLE:")
        for model in (m.Chapter, m.Topic, m.Concept, m.ConceptPrerequisite, m.SourceChunk):
            assert db.scalar(select(func.count()).select_from(model)) == 0


@pytest.mark.parametrize("bad", ["source", "prerequisite", "cycle"])
def test_invalid_outline_rolls_back_even_if_caller_commits(uploaded, factory, bad):
    class InvalidOutline(OutlineModel):
        async def generate_structured(self, messages, schema):
            result = await super().generate_structured(messages, schema)
            concepts = result.chapters[0].topics[0].concepts
            if bad == "source":
                concepts[1].source_numbers = [999]
            elif bad == "prerequisite":
                concepts[1].prerequisite_names = ["Invented"]
            else:
                concepts[0].prerequisite_names = ["Root uptake"]
            return result

    course_id, _, storage = uploaded
    with factory() as db:
        with pytest.raises(AppError) as error:
            run(service(db, storage, model=InvalidOutline()).process_course(course_id))
        assert error.value.code == "INVALID_KNOWLEDGE"
        db.commit()
        assert db.scalar(select(func.count()).select_from(m.Chapter)) == 0


def test_actual_progress_phases(uploaded, factory):
    course_id, _, storage = uploaded
    stages = []
    token = observer.set(lambda stage, progress: stages.append((stage, progress)))
    try:
        with factory() as db:
            run(service(db, storage).process_course(course_id))
    finally:
        observer.reset(token)
    assert stages == [
        ("extracting_content", 10),
        ("analyzing_content", 30),
        ("building_course", 50),
        ("creating_embeddings", 75),
        ("indexing", 90),
    ]


def test_pgvector_query_compilation():
    course_id, concept_id = uuid4(), uuid4()
    query = vector_search_statement(course_id, [1.0] + [0.0] * 1023, 3, concept_id)
    compiled = query.compile(dialect=postgresql.dialect())
    sql = str(compiled)
    assert "<=>" in sql and "ORDER BY distance" in sql and "LIMIT" in sql
    assert "source_chunks.course_id =" in sql and "source_chunks.concept_id =" in sql
    assert "EXISTS" in sql and "source_chunk_concepts.concept_id =" in sql
    assert course_id in compiled.params.values() and concept_id in compiled.params.values()


def test_no_config_boots_and_real_embedding_fails_explicitly(client, factory, monkeypatch):
    settings = Settings(
        _env_file=None,
        knowledge_provider="questoriaa",
        embedding_provider="openai",
        embedding_base_url="",
        embedding_api_key="",
    )
    monkeypatch.setattr(factories, "get_settings", lambda: settings)
    assert client.get("/openapi.json").status_code == 200
    with factory() as db:
        assert isinstance(factories.create_knowledge(db), QuestoriaaKnowledgeService)
    with pytest.raises(AppError) as error:
        run(factories.create_embeddings().embed_query("question"))
    assert error.value.code == "EMBEDDING_UNAVAILABLE"


def test_embedding_endpoint_batching_order_and_query_instruction(monkeypatch):
    requests = []

    def handler(request):
        payload = json.loads(request.content)
        requests.append(payload)
        return httpx.Response(
            200,
            json={
                "data": [
                    {"index": i, "embedding": [float(i + 1)] + [0.0] * 1023}
                    for i in reversed(range(len(payload["input"])))
                ]
            },
        )

    client_type = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: client_type(transport=httpx.MockTransport(handler), **kwargs),
    )
    provider = OpenAIEmbeddingProvider("https://embedding.test/v1", "test-only", "qwen")
    assert len(run(provider.embed_texts(["text"] * 17))) == 17
    assert [len(request["input"]) for request in requests] == [16, 1]
    assert run(provider.embed_query("question"))[0] == 1.0
    assert requests[-1]["input"][0].startswith("Instruct:")
    assert requests[-1]["input"][0].endswith("Query: question")


@pytest.mark.parametrize("payload", [{"data": []}, {"data": [{"index": 0, "embedding": [1]}]}])
def test_embedding_endpoint_invalid_responses(monkeypatch, payload):
    client_type = httpx.AsyncClient
    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: client_type(
            transport=httpx.MockTransport(lambda request: httpx.Response(200, json=payload)),
            **kwargs,
        ),
    )
    with pytest.raises(AppError) as error:
        run(OpenAIEmbeddingProvider("https://embedding.test/v1", "", "qwen").embed_texts(["text"]))
    assert error.value.code == "EMBEDDING_UNAVAILABLE"


def test_database_survives_engine_recreation(tmp_path):
    url = f"sqlite:///{(tmp_path / 'persistent.db').as_posix()}"
    engine = make_engine(url)
    Base.metadata.create_all(engine)
    storage = LocalStorageService(tmp_path / "storage", "/api/v1")
    data = pdf_bytes()
    storage.root.mkdir()
    storage.path("notes.pdf").write_bytes(data)
    with sessionmaker(engine)() as db:
        user = m.User(clerk_user_id="persistent-test")
        db.add(user)
        db.flush()
        course = m.Course(user_id=user.id, title="Botany")
        db.add(course)
        db.flush()
        course_id = course.id
        db.add(
            m.CourseFile(
                course_id=course_id,
                filename="notes.pdf",
                source_type="pdf",
                object_key="notes.pdf",
                content_type="application/pdf",
                size_bytes=len(data),
                upload_completed=True,
            )
        )
        db.commit()
        run(service(db, storage).process_course(course_id))
        db.commit()
    engine.dispose()

    engine = make_engine(url)
    with sessionmaker(engine)() as db:
        fresh = service(db, storage)
        assert len(run(fresh.get_course_map(course_id)).chapters) == 1
        found = run(fresh.retrieve_context(course_id, "water soil"))
        assert len(found) == 2 and found[0].source.page_number == 2
    engine.dispose()


def test_missing_embeddings_are_excluded_from_semantic_search(uploaded, factory):
    course_id, file_id, _ = uploaded
    with factory() as db:
        db.add(
            m.SourceChunk(
                course_id=course_id,
                file_id=file_id,
                content="Legacy unembedded text",
                embedding=None,
            )
        )
        db.commit()
        store = DatabaseChunkStore(db)
        assert store.search(course_id, run(MockEmbeddingProvider().embed_query("text"))) == []
        chunk_id = db.scalar(select(m.SourceChunk.id))
        assert len(store.get_by_ids(course_id, [chunk_id])) == 1


def test_store_rejects_cross_course_file(uploaded, factory):
    course_id, file_id, _ = uploaded
    with factory() as db:
        chunk = m.SourceChunk(
            course_id=uuid4(),
            file_id=file_id,
            content="Wrong course",
            embedding=[1.0] + [0.0] * 1023,
        )
        with pytest.raises(ValueError, match="file must belong"):
            DatabaseChunkStore(db).add([chunk])


def test_worker_reports_unsupported_input_without_graph(uploaded, factory):
    course_id, file_id, storage = uploaded
    with factory() as db:
        db.get(m.CourseFile, file_id).source_type = "image"
        job = m.ProcessingJob(course_id=course_id, file_id=file_id)
        db.add(job)
        db.commit()
        run(execute_job(db, job.id, service(db, storage)))
        assert job.status == "failed" and job.error.startswith("UNSUPPORTED_INGESTION:")
        assert db.scalar(select(func.count()).select_from(m.SourceChunk)) == 0


def test_mock_outline_supports_multiple_topics():
    sources = [{"number": i, "content": f"Source content {i}"} for i in range(1, 22)]
    outline = run(
        MockModelProvider().generate_structured(
            [{"role": "user", "content": json.dumps({"title": "Example", "sources": sources})}],
            CourseDraft,
        )
    )
    assert [len(t.concepts) for t in outline.chapters[0].topics] == [20, 1]


class SharedOutlineModel(MockModelProvider):
    async def generate_structured(self, messages, schema):
        if schema is not CourseDraft:
            return await super().generate_structured(messages, schema)
        return CourseDraft.model_validate(
            {
                "chapters": [
                    {
                        "title": "Shared evidence",
                        "topics": [
                            {
                                "title": "Two concepts",
                                "concepts": [
                                    {
                                        "name": "Concept A",
                                        "description": "First interpretation.",
                                        "source_numbers": [1, 1],
                                    },
                                    {
                                        "name": "Concept B",
                                        "description": "Second interpretation.",
                                        "source_numbers": [1],
                                    },
                                    {
                                        "name": "Concept B",
                                        "description": "Duplicate citation.",
                                        "source_numbers": [1],
                                    },
                                ],
                            }
                        ],
                    }
                ]
            }
        )


@pytest.fixture
def shared_evidence(uploaded, factory, monkeypatch):
    course_id, file_id, storage = uploaded
    data = pdf_bytes(["Plants use sunlight and water to produce sugar."])
    with factory() as db:
        file = db.get(m.CourseFile, file_id)
        storage.path(file.object_key).write_bytes(data)
        file.size_bytes = len(data)
        db.commit()
        run(service(db, storage, model=SharedOutlineModel()).process_course(course_id))
        db.commit()
        chunk = db.scalar(select(m.SourceChunk))
        concepts = {c.name: c.id for c in db.scalars(select(m.Concept))}
        chunk_id = chunk.id
    settings = Settings(
        _env_file=None,
        environment="test",
        knowledge_provider="questoriaa",
        embedding_provider="mock",
    )
    monkeypatch.setattr(factories, "get_settings", lambda: settings)
    monkeypatch.setattr(factories, "get_storage", lambda: storage)
    monkeypatch.setattr(factories, "create_model", SharedOutlineModel)
    return course_id, chunk_id, concepts


def test_shared_evidence_retrieval_and_adaptive_learning(shared_evidence, factory, client):
    course_id, chunk_id, concepts = shared_evidence
    with factory() as db:
        associations = db.execute(
            select(m.SourceChunkConcept.chunk_id, m.SourceChunkConcept.concept_id)
        ).all()
        assert set(associations) == {(chunk_id, id_) for id_ in concepts.values()}
        knowledge = factories.create_knowledge(db)
        for concept_id in concepts.values():
            semantic = run(knowledge.retrieve_context(course_id, "sunlight", concept_id))
            exact = run(knowledge.get_context_by_ids(course_id, [chunk_id], concept_id))
            assert [chunk.chunk_id for chunk in semantic] == [chunk_id]
            assert [chunk.chunk_id for chunk in exact] == [chunk_id]
            assert semantic[0].concept_id == exact[0].concept_id == concept_id
        assert db.get(m.SourceChunk, chunk_id).concept_id == concepts["Concept A"]
        assert [
            chunk.chunk_id for chunk in run(knowledge.retrieve_context(course_id, "sunlight"))
        ] == [chunk_id]
        assert run(knowledge.get_context_by_ids(uuid4(), [chunk_id], concepts["Concept B"])) == []
        # The legacy adapter obeys the same shared-evidence scope.
        mock_context = run(
            MockKnowledgeService(db).retrieve_context(course_id, "sunlight", concepts["Concept B"])
        )
        assert mock_context[0].chunk_id == chunk_id
        assert mock_context[0].concept_id == concepts["Concept B"]
    response = client.post(
        "/api/v1/challenges/next",
        json={"course_id": str(course_id), "concept_id": str(concepts["Concept B"])},
    )
    assert response.status_code == 201, response.text
    challenge = response.json()
    assert challenge["source_chunk_ids"] == [str(chunk_id)]
    response = client.post(
        "/api/v1/attempts",
        json={
            "challenge_id": challenge["id"],
            "answer": {"selected_option": 0},
            "time_taken_seconds": 5,
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["correct"] is True
    assert response.json()["sources"][0]["page_number"] == 1
    with factory() as db:
        assert db.get(m.SourceChunk, chunk_id).concept_id == concepts["Concept A"]
        assert (
            db.scalar(
                select(m.MasteryState).where(m.MasteryState.concept_id == concepts["Concept B"])
            ).attempt_count
            == 1
        )


@pytest.mark.parametrize("target", ["chunk", "primary_concept", "secondary_concept", "course"])
def test_shared_evidence_cascades(shared_evidence, factory, target):
    course_id, chunk_id, concepts = shared_evidence
    with factory() as db:
        if target == "chunk":
            db.delete(db.get(m.SourceChunk, chunk_id))
        elif target == "course":
            db.delete(db.get(m.Course, course_id))
        else:
            name = "Concept A" if target == "primary_concept" else "Concept B"
            db.delete(db.get(m.Concept, concepts[name]))
        db.commit()
        links = db.execute(
            select(m.SourceChunkConcept.chunk_id, m.SourceChunkConcept.concept_id)
        ).all()
        if target in {"chunk", "course"}:
            assert links == []
        else:
            remaining = (
                concepts["Concept B"] if target == "primary_concept" else concepts["Concept A"]
            )
            assert links == [(chunk_id, remaining)]
            assert (
                DatabaseChunkStore(db).get_by_ids(course_id, [chunk_id], remaining)[0].concept_id
                == remaining
            )
            if target == "primary_concept":
                assert db.get(m.SourceChunk, chunk_id).concept_id is None


def test_shared_evidence_pair_is_unique(shared_evidence, factory):
    _, chunk_id, concepts = shared_evidence
    with factory() as db:
        with pytest.raises(IntegrityError), db.begin_nested():
            db.add(m.SourceChunkConcept(chunk_id=chunk_id, concept_id=concepts["Concept B"]))
            db.flush()
        assert db.scalar(select(func.count()).select_from(m.SourceChunkConcept)) == 2
