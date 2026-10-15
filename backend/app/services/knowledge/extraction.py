"""PDF/PPTX text extraction adapted from the A/AJ extractors.

Never merge different pages/slides: our public citation has a single location.
"""

import re
from dataclasses import dataclass
from io import BytesIO

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.core.errors import AppError
from app.services.knowledge.chunking import chunk_text


@dataclass(frozen=True)
class ExtractedChunk:
    content: str
    page_number: int | None = None
    slide_number: int | None = None


def clean_text(text: str) -> str:
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text.replace("\r", "\n"))
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()


def shape_lines(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [line for child in shape.shapes for line in shape_lines(child)]
    lines = []
    if shape.has_text_frame:
        lines.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    if shape.has_table:
        lines.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for row in shape.table.rows
            if any(cell.text.strip() for cell in row.cells)
        )
    return lines


def extract(data: bytes, source_type: str) -> list[ExtractedChunk]:
    if source_type not in {"pdf", "pptx"}:
        raise AppError(
            422, "UNSUPPORTED_INGESTION", "Knowledge ingestion supports PDF/PPTX text only."
        )
    try:
        chunks = []
        if source_type == "pdf":
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.needs_pass:
                    raise AppError(
                        422, "INVALID_DOCUMENT", "Encrypted PDF requires an unlocked upload."
                    )
                for number, page in enumerate(document, 1):
                    text = clean_text(page.get_text("text"))
                    if not text:
                        continue
                    chunks.extend(
                        ExtractedChunk(piece, page_number=number) for piece in chunk_text(text)
                    )
        else:
            presentation = Presentation(BytesIO(data))
            for number, slide in enumerate(presentation.slides, 1):
                lines = [line for shape in slide.shapes for line in shape_lines(shape)]
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        lines.append(f"Speaker notes: {notes}")
                text = clean_text("\n".join(lines))
                if not text:
                    continue
                chunks.extend(
                    ExtractedChunk(piece, slide_number=number) for piece in chunk_text(text)
                )
        if not chunks:
            raise AppError(422, "TEXTLESS_DOCUMENT", "No extractable text was found.")
        return chunks
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            422, "INVALID_DOCUMENT", "The uploaded document could not be parsed."
        ) from exc
